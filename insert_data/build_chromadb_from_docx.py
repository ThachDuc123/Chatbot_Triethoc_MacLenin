"""Build ChromaDB directly from DOCX philosophy textbooks.

Why:
- User wants the bot to "learn" from DOCX (giáo trình) instead of CSV.
- In practice this is RAG indexing: parse DOCX -> chunk -> embed -> store in Chroma.

This script reuses the existing embedding backend convention:
collection_name = {base_model_name}__{embedding_backend}
so query and indexing stay consistent and avoid dimension mismatch.
"""

from __future__ import annotations

import argparse
import os
import re
from pathlib import Path
from typing import Dict, List

import chromadb


def _clean_text(text: str) -> str:
    text = re.sub(r"\s+", " ", text or " ")
    text = re.sub(r"[\x00-\x08\x0b-\x0c\x0e-\x1f\x7f-\x9f]", "", text)
    return text.strip()


def _split_into_chunks(text: str, chunk_size: int, overlap: int) -> List[str]:
    chunks: List[str] = []
    start = 0
    text_length = len(text)

    while start < text_length:
        end = start + chunk_size
        chunk = text[start:end]

        if end < text_length:
            for delimiter in [". ", ".\n", "; ", "! ", "? ", "\n\n", "\n"]:
                last_delim = chunk.rfind(delimiter)
                if last_delim != -1 and last_delim > int(chunk_size * 0.8):
                    chunk = chunk[: last_delim + len(delimiter)]
                    end = start + len(chunk)
                    break

        chunk = chunk.strip()
        if len(chunk) >= 50:
            chunks.append(chunk)

        start = end - overlap
        if start >= text_length - overlap:
            break

    return chunks


def _read_docx_text(docx_path: Path) -> str:
    try:
        from docx import Document
    except Exception as e:  # pragma: no cover
        raise ImportError("Missing dependency: python-docx. Install it to read .docx") from e

    doc = Document(str(docx_path))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
    return "\n".join(paragraphs)


def _read_docx_paragraphs(docx_path: Path) -> List[str]:
    """Read DOCX and return a list of clean paragraphs (no empty)."""
    try:
        from docx import Document
    except Exception as e:  # pragma: no cover
        raise ImportError("Missing dependency: python-docx. Install it to read .docx") from e

    doc = Document(str(docx_path))
    return [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]


def _chunks_from_paragraphs(paragraphs: List[str], chunk_size: int, overlap: int) -> List[str]:
    """Chunk by accumulating paragraphs until reaching chunk_size.

    This preserves section boundaries better than pure character slicing.
    """
    chunks: List[str] = []
    cur: List[str] = []
    cur_len = 0
    for para in paragraphs:
        p = _clean_text(para)
        if not p:
            continue

        # If a single paragraph is huge, fall back to char splitting.
        if len(p) > chunk_size:
            if cur:
                chunks.append("\n".join(cur).strip())
                cur, cur_len = [], 0
            chunks.extend(_split_into_chunks(p, chunk_size=chunk_size, overlap=overlap))
            continue

        if cur_len + len(p) + 1 > chunk_size and cur:
            chunks.append("\n".join(cur).strip())
            # overlap: keep last N chars worth of text
            if overlap > 0:
                tail = "\n".join(cur)[-overlap:]
                cur = [tail] if tail.strip() else []
                cur_len = len(tail)
            else:
                cur, cur_len = [], 0

        cur.append(p)
        cur_len += len(p) + 1

    if cur:
        chunks.append("\n".join(cur).strip())
    return [c for c in chunks if len(c) >= 50]


def _get_collection_name(model_name: str, embedding_backend: str) -> str:
    base_name = model_name.split("/")[-1] if "/" in model_name else model_name
    return f"{base_name}__{embedding_backend}"


def _embed_texts(texts: List[str], model_name: str, embedding_backend: str) -> List[List[float]]:
    if embedding_backend == "fastembed":
        try:
            from fastembed import TextEmbedding
        except Exception as e:  # pragma: no cover
            raise ImportError("fastembed is not installed. Install it or use sentence_transformers") from e

        fe = TextEmbedding(name=model_name, max_length=512)
        return [v.tolist() for v in fe.embed(texts)]

    # default: sentence_transformers
    from sentence_transformers import SentenceTransformer

    model = SentenceTransformer(model_name, trust_remote_code=True)
    return [model.encode(t).tolist() for t in texts]


def build_from_docx(
    docx_path: str,
    persist_dir: str = "./chroma_db",
    model_name: str = "Alibaba-NLP/gte-multilingual-base",
    embedding_backend: str = "fastembed",
    chunk_size: int = 800,
    overlap: int = 100,
    source_name: str | None = None,
) -> None:
    path = Path(docx_path)
    if not path.exists():
        raise FileNotFoundError(f"DOCX not found: {path}")

    paragraphs = _read_docx_paragraphs(path)
    chunks = _chunks_from_paragraphs(paragraphs, chunk_size=chunk_size, overlap=overlap)
    if not chunks:
        raise RuntimeError("No chunks produced from DOCX (document may be empty)")

    src = source_name or path.stem
    texts = []
    previews = []
    for c in chunks:
        prev = c.replace("\n", " ")[:80]
        previews.append(prev)
        texts.append(
            f"Tiêu đề: Giáo trình Triết học Mác-Lênin, Nguồn: {src}, Trích đoạn: {prev}, Nội dung: {c}"
        )

    embeddings = _embed_texts(texts, model_name=model_name, embedding_backend=embedding_backend)

    client = chromadb.PersistentClient(path=persist_dir)
    collection_name = _get_collection_name(model_name=model_name, embedding_backend=embedding_backend)
    collection = client.get_or_create_collection(name=collection_name)

    ids = [f"{(source_name or path.stem)}_chunk_{i}" for i in range(len(texts))]
    metadatas: List[Dict[str, str]] = []
    for i in range(len(texts)):
        metadatas.append(
            {
                "source_file": str(src),
                "chunk_index": str(i),
                "category": "Triết học",
                "file_path": str(path),
                "preview": previews[i],
            }
        )

    # Upsert to allow re-indexing without deleting the collection.
    collection.upsert(ids=ids, documents=texts, embeddings=embeddings, metadatas=metadatas)
    print(f"✅ Upserted {len(texts)} chunks into collection `{collection_name}`")


def main() -> None:
    parser = argparse.ArgumentParser(description="Build ChromaDB from a DOCX textbook")
    parser.add_argument("--docx_path", type=str, required=True, help="Path to .docx")
    parser.add_argument("--persist_dir", type=str, default="./chroma_db")
    parser.add_argument("--model_name", type=str, default="Alibaba-NLP/gte-multilingual-base")
    parser.add_argument(
        "--embedding_backend",
        type=str,
        default="fastembed",
        choices=["fastembed", "sentence_transformers"],
    )
    parser.add_argument("--chunk_size", type=int, default=800)
    parser.add_argument("--overlap", type=int, default=100)
    parser.add_argument("--source_name", type=str, default=None)
    args = parser.parse_args()

    build_from_docx(
        docx_path=args.docx_path,
        persist_dir=args.persist_dir,
        model_name=args.model_name,
        embedding_backend=args.embedding_backend,
        chunk_size=args.chunk_size,
        overlap=args.overlap,
        source_name=args.source_name,
    )


if __name__ == "__main__":
    main()
