"""
Script siêu nhanh - xử lý giáo trình trong 30 giây
"""

print("🚀 Bắt đầu xử lý siêu nhanh...")

from docx import Document
import pandas as pd
from pathlib import Path

# Đọc file
print("📖 Đọc file...")
doc = Document("philosophy_books/1. Giao trinh Triet hoc Mac - Lenin 2021.docx")

# Lấy text nhanh - không làm sạch phức tạp
print(f"✅ Đã đọc {len(doc.paragraphs)} đoạn")
print("📝 Trích xuất text...")

text_parts = [p.text for p in doc.paragraphs if p.text.strip()]
full_text = " ".join(text_parts)

print(f"✅ Tổng {len(full_text):,} ký tự")

# Chia chunks đơn giản
print("✂️  Chia chunks...")
chunk_size = 1000
chunks = [full_text[i:i+chunk_size] for i in range(0, len(full_text), chunk_size-100)]
chunks = [c for c in chunks if len(c) >= 200]

print(f"✅ Tạo {len(chunks)} chunks")

# Tạo data
print("📊 Tạo DataFrame...")
data = [{
    '_id': f'triet_{i}',
    'title': f'Triết học Mác-Lênin (P{i+1})',
    'content': c,
    'source_file': 'GiaoTrinhTrietHoc',
    'chunk_index': i,
    'category': 'Triết học',
    'combined_information': f'Tiêu đề: Triết học Mác-Lênin (P{i+1}), Danh mục: Triết học, Nội dung: {c}'
} for i, c in enumerate(chunks)]

df = pd.DataFrame(data)

# Lưu
print("💾 Lưu file...")
Path("data").mkdir(exist_ok=True)
df.to_csv("data/philosophy_data.csv", index=False, encoding='utf-8-sig')

print("\n" + "="*50)
print("✅ HOÀN THÀNH!")
print("="*50)
print(f"📁 File: data/philosophy_data.csv")
print(f"📊 Chunks: {len(df)}")
print(f"📝 Ký tự: {len(full_text):,}")
print("\n🎉 Chạy chatbot:")
print("python serve.py --mode online --model_name gemini --model_version gemini-1.5-flash")
print("="*50)

# Hiện mẫu
print("\n📚 Chunk đầu tiên:")
print(data[0]['content'][:300] + "...")
