"""
Script xử lý TOÀN BỘ giáo trình triết học - phiên bản tối ưu
"""

from docx import Document
import pandas as pd
import re
from pathlib import Path
import time

def process_full_document():
    print("="*60)
    print("🎓 XỬ LÝ GIÁO TRÌNH TRIẾT HỌC MÁC-LÊNIN")
    print("="*60)
    
    # Đọc file
    file_path = "philosophy_books/1. Giao trinh Triet hoc Mac - Lenin 2021.docx"
    print(f"\n📖 Đang đọc file: {file_path}")
    print("⏳ Vui lòng đợi...")
    
    start_time = time.time()
    
    doc = Document(file_path)
    
    # Lấy tất cả văn bản
    print(f"✅ Đã mở file ({len(doc.paragraphs)} đoạn văn)")
    print("📝 Đang trích xuất nội dung...")
    
    paragraphs = []
    for i, para in enumerate(doc.paragraphs):
        if i % 100 == 0:
            print(f"   → Đã xử lý {i}/{len(doc.paragraphs)} đoạn...")
        if para.text.strip():
            paragraphs.append(para.text.strip())
    
    text = " ".join(paragraphs)
    
    print(f"✅ Đã trích xuất {len(text):,} ký tự")
    
    # Làm sạch
    print("🧹 Đang làm sạch văn bản...")
    text = re.sub(r'\s+', ' ', text)
    text = re.sub(r'[\x00-\x08\x0b-\x0c\x0e-\x1f\x7f-\x9f]', '', text)
    text = text.strip()
    
    print(f"✅ Sau khi làm sạch: {len(text):,} ký tự")
    
    # Chia thành chunks
    print("✂️  Đang chia thành các chunks...")
    chunk_size = 800  # Tăng lên để giảm số chunks
    overlap = 100
    chunks = []
    
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunk = text[start:end]
        
        # Cắt tại dấu câu nếu có thể
        if end < len(text):
            for delimiter in ['. ', '.\n', '; ', '! ', '? ', '\n']:
                last_delim = chunk.rfind(delimiter)
                if last_delim > chunk_size * 0.8:  # Chỉ cắt nếu gần cuối
                    chunk = chunk[:last_delim + len(delimiter)]
                    end = start + len(chunk)
                    break
        
        if len(chunk.strip()) >= 100:  # Chỉ giữ chunks đủ dài
            chunks.append(chunk.strip())
        
        start = end - overlap
        
        if start >= len(text) - overlap:
            break
    
    print(f"✅ Đã tạo {len(chunks)} chunks")
    
    # Tạo DataFrame
    print("📊 Đang tạo DataFrame...")
    data = []
    
    for idx, chunk in enumerate(chunks):
        if idx % 100 == 0:
            print(f"   → Đã xử lý {idx}/{len(chunks)} chunks...")
        
        # Trích xuất tiêu đề từ chunk (50 ký tự đầu)
        title_preview = chunk[:50].replace('\n', ' ')
        if len(title_preview) < 50 and len(chunk) > 50:
            title_preview = chunk[:50]
        
        data.append({
            '_id': f'triet_hoc_chunk_{idx:04d}',
            'title': f'Giáo trình Triết học Mác-Lênin (Phần {idx+1})',
            'content': chunk,
            'source_file': 'Giao_trinh_Triet_hoc_Mac_Lenin_2021',
            'chunk_index': idx,
            'category': 'Triết học',
            'combined_information': f'Tiêu đề: Giáo trình Triết học Mác-Lênin (Phần {idx+1}), Danh mục: Triết học, Nguồn: Giáo trình chính thức năm 2021, Nội dung: {chunk}'
        })
    
    df = pd.DataFrame(data)
    
    # Lưu CSV
    print("💾 Đang lưu vào CSV...")
    Path("data").mkdir(exist_ok=True)
    output_file = "data/philosophy_data.csv"
    df.to_csv(output_file, index=False, encoding='utf-8-sig')
    
    elapsed_time = time.time() - start_time
    
    print("\n" + "="*60)
    print("✅ HOÀN THÀNH!")
    print("="*60)
    print(f"📁 File đầu ra: {output_file}")
    print(f"📊 Số lượng chunks: {len(df)}")
    print(f"📝 Tổng số ký tự: {sum(len(d['content']) for d in data):,}")
    print(f"⏱️  Thời gian xử lý: {elapsed_time:.1f} giây")
    print("\n🎉 Bây giờ bạn có thể chạy chatbot:")
    print("   python serve.py --mode online --model_name gemini --model_version gemini-1.5-flash")
    print("="*60)
    
    # Hiển thị vài chunks mẫu
    print("\n📚 Xem thử 3 chunks đầu tiên:\n")
    for i in range(min(3, len(data))):
        print(f"--- Chunk {i+1}: {data[i]['title']} ---")
        print(data[i]['content'][:250] + "...\n")

if __name__ == "__main__":
    try:
        process_full_document()
    except KeyboardInterrupt:
        print("\n\n⚠️  Đã dừng bởi người dùng")
    except Exception as e:
        print(f"\n\n❌ Lỗi: {e}")
        import traceback
        traceback.print_exc()
